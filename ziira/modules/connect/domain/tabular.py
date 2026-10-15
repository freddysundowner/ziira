"""Any uploaded file as a plain grid of text: CSV, Excel (.xlsx), a text PDF, or (through the OCR step) a picture.

`load` looks at what the file *is* (its first bytes), never at its name, and returns a `RawTable`: lines of cells with
the line number, the page, and where the text came from. Finding the header, the columns and the formats is the job of
`layout`; this module only reads. Problems raise `FileProblem` with a message a shop owner can act on; two subclasses
tell the caller what to do next (`PasswordNeeded`, `NeedsOcr`).
"""

import io
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal

from ziira.modules.connect.domain import csvfile
from ziira.modules.connect.domain.csvfile import FileProblem

MAX_ROWS = csvfile.MAX_ROWS
MAX_UNZIPPED = 150 * 1024 * 1024
PNG, JPEG = b"\x89PNG\r\n\x1a\n", b"\xff\xd8\xff"
IMAGE_MAGIC = (PNG, JPEG, b"II*\x00", b"MM\x00*")
OLD_EXCEL = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class PasswordNeeded(FileProblem):
    """The PDF is protected. `wrong` is True when a password was given and did not open it."""

    def __init__(self, wrong: bool) -> None:
        super().__init__("The password does not open this file. Check it and try again." if wrong
                         else "This PDF is protected with a password. Type the password to open it.")
        self.wrong = wrong


class NeedsOcr(FileProblem):
    """The file is a picture (a photo, or a PDF made of scanned pages): its text must be read by the OCR engine."""

    def __init__(self, kind: str, pages: int) -> None:
        super().__init__("This file is a picture, not text: it has to be read from the image.")
        self.kind, self.pages = kind, pages


@dataclass
class RawRow:
    """One line of the file: cells as text. `number` is the line (CSV), row (Excel) or counted line (PDF)."""

    number: int
    cells: list[str]
    page: int = 0
    low_confidence: bool = False        # read by OCR with a doubtful word in it


@dataclass
class RawTable:
    rows: list[RawRow]
    source: str                          # csv | excel | pdf | scan
    warnings: list[str] = field(default_factory=list)
    delimiter: str | None = None
    ocr: bool = False


def sniff(data: bytes) -> str:
    """csv | excel | old_excel | pdf | image: decided from the first bytes."""
    head = data[:16]
    if head.startswith(b"%PDF") or b"%PDF-" in data[:1024]:
        return "pdf"
    if head.startswith(b"PK\x03\x04"):
        return "excel"
    if head.startswith(OLD_EXCEL):
        return "old_excel"
    if head.startswith(IMAGE_MAGIC) or (head.startswith(b"RIFF") and data[8:12] == b"WEBP"):
        return "image"
    return "csv"


def _cell(value: object) -> str:
    """An Excel cell as the text a CSV would hold: dates as ISO, numbers without float noise."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d") if value.time() == time(0) else value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float):
        return format(Decimal(repr(value)), "f")
    return " ".join(str(value).split()) if isinstance(value, str) else str(value)


def read_excel(data: bytes) -> RawTable:
    """The sheet with the most filled rows of an .xlsx workbook (formulas read as last saved)."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_UNZIPPED:
                raise FileProblem("This Excel file is too large to read. Save it again with fewer rows.")
            if "xl/workbook.xml" not in archive.namelist():
                raise FileProblem("This file is not an Excel workbook (.xlsx).")
    except zipfile.BadZipFile:
        raise FileProblem("This Excel file could not be opened. Save it again as .xlsx or CSV.") from None
    from openpyxl import load_workbook  # imported late: only Excel uploads need it
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:                                                              # noqa: BLE001 - any broken file
        raise FileProblem("This Excel file could not be opened. If it asks for a password, save a copy without one, "
                          "or download the statement as CSV.") from None
    best: list[RawRow] = []
    best_name = ""
    try:
        for sheet in book.worksheets:
            if getattr(sheet, "sheet_state", "visible") != "visible":
                continue
            rows: list[RawRow] = []
            for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                cells = [_cell(v) for v in values]
                if any(c.strip() for c in cells):
                    rows.append(RawRow(number, cells))
                if len(rows) > MAX_ROWS:
                    raise FileProblem(f"The file has more than {MAX_ROWS} rows. Split it into smaller files.")
            if len(rows) > len(best):
                best, best_name = rows, sheet.title
    finally:
        book.close()
    if not best:
        raise FileProblem("The Excel file has no data.")
    warnings = [f"The sheet '{best_name}' was read."] if len(book.worksheets) > 1 else []
    return RawTable(best, "excel", warnings)


def read_csv(data: bytes, delimiter: str | None) -> RawTable:
    rows, chosen, warnings = csvfile.raw_rows(data, delimiter=delimiter)
    for number, cells in rows:
        if any(len(c) > csvfile.MAX_CELL for c in cells):
            raise FileProblem(f"Row {number} has a value longer than {csvfile.MAX_CELL} characters.")
    return RawTable([RawRow(n, [c.strip() for c in cells]) for n, cells in rows], "csv", warnings, chosen)


def load(data: bytes, *, vocabulary: set[str], password: str | None = None, delimiter: str | None = None) -> RawTable:
    """Read an uploaded file of any supported kind. `vocabulary` is the set of normalised column names that mark a
    header line (used to find the table inside a PDF). Raises FileProblem, PasswordNeeded or NeedsOcr."""
    if len(data) > csvfile.MAX_BYTES:
        raise FileProblem(f"The file is larger than {csvfile.MAX_BYTES // (1024 * 1024)} MB. "
                          "Split it into smaller files.")
    if not data.strip():
        raise FileProblem("The file is empty.")
    kind = sniff(data)
    if kind == "old_excel":
        raise FileProblem("This is an old Excel file (.xls). Open it in Excel and use Save As > Excel Workbook "
                          "(.xlsx), "
                          "or CSV.")
    if kind == "excel":
        return read_excel(data)
    if kind == "pdf":
        from ziira.modules.connect.domain import pdfreader  # imported late: the PDF libraries are heavy
        return pdfreader.read_pdf(data, vocabulary=vocabulary, password=password)
    if kind == "image":
        raise NeedsOcr("image", 1)
    return read_csv(data, delimiter)
